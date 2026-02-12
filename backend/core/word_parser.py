from docx import Document
import base64


def extract_word_spec(word_path: str):

    doc = Document(word_path)

    # Extract image
    template_image_b64 = None
    for rel in doc.part._rels:
        rel_obj = doc.part._rels[rel]
        if "image" in rel_obj.target_ref:
            image_data = rel_obj.target_part.blob
            template_image_b64 = base64.b64encode(image_data).decode()
            break

    # Extract table mapping
    mapping = []
    if doc.tables:
        table = doc.tables[0]
        for row in table.rows[1:]:
            position = int(row.cells[0].text.strip())
            json_path = row.cells[2].text.strip()

            mapping.append({
                "position": position,
                "json_path": json_path
            })

    return template_image_b64, mapping